"""
Data Migration Script: Import ALL historical production data from Excel into SQLite.

Reads:
- Produtividade_Geral.xlsx (consolidated data, sector tabs, manual line)
- All individual Produtividade*.xlsx files from EXCEL_FOLDER

Populates:
- producao_diaria: automatic line daily totals (2024 to present), turno, absenteeism
- producao_manual: detailed manual packaging records
- producao_manual_totais: daily aggregated manual totals
- observacoes_diarias: sector daily observations & machinery downtime justifications
"""

import os
import sys
import glob
from datetime import datetime, date
from collections import defaultdict
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from config import EXCEL_FOLDER, DASHBOARD_DB
from database import init_db, get_db

# Turno alternation reference: 2026-09-30 = Caio. Even days = Caio, odd = Douglas.
TURNO_REF = date(2026, 9, 30)


def get_turno(data_str_or_date):
    """Calculate turno based on alternating schedule from reference date."""
    if isinstance(data_str_or_date, (datetime, date)):
        dt = data_str_or_date
        if isinstance(dt, datetime):
            dt = dt.date()
    else:
        dt = datetime.strptime(str(data_str_or_date)[:10], '%Y-%m-%d').date()
    delta = (dt - TURNO_REF).days
    return 'Caio' if delta % 2 == 0 else 'Douglas'


def parse_date(val):
    """Robust date parser handling datetime, date, and various string formats."""
    if val is None:
        return None
    if isinstance(val, datetime):
        dt = val.date()
    elif isinstance(val, date):
        dt = val
    elif isinstance(val, str):
        val = val.strip()
        for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%Y-%m-%d %H:%M:%S', '%d-%m-%Y'):
            try:
                dt = datetime.strptime(val, fmt).date()
                break
            except ValueError:
                pass
        else:
            return None
    else:
        return None

    # Sanitize known Excel typos (e.g. 2055 -> 2025)
    if dt.year == 2055:
        dt = dt.replace(year=2025)

    if not (2020 <= dt.year <= 2030):
        return None

    return dt.strftime('%Y-%m-%d')


def parse_int(val):
    """Safely convert cell value to integer."""
    if val is None:
        return 0
    try:
        s = str(val).strip().replace('-', '0')
        return int(float(s))
    except (ValueError, TypeError):
        return 0


def clean_obs(val):
    """Sanitize observation and machine problem text."""
    if val is None:
        return None
    s = str(val).strip()
    if not s or s.lower() in ('none', 'nan', 'null', '-', ''):
        return None
    if s.lower() in ('envase 01 e 02', 'envase 1 e 2', 'envase 01 e 02 ', 'envase 1 e 2 '):
        return None
    return s


def find_sheet_ci(wb, target_name):
    """Case-insensitive sheet lookup."""
    target_clean = target_name.lower().strip()
    for s in wb.sheetnames:
        if s.lower().strip() == target_clean:
            return s
    return None


def find_consolidado_cols(ws):
    """Dynamically find column indices in Consolidado tab."""
    for r in range(1, 10):
        row_vals = [ws.cell(r, c).value for c in range(1, min(ws.max_column + 1, 30))]
        row_str = [str(v).strip().lower() if v is not None else '' for v in row_vals]
        if any('data' in s for s in row_str):
            col_map = {}
            for idx, s in enumerate(row_str):
                col_num = idx + 1
                if 'data' in s and 'date' not in col_map:
                    col_map['date'] = col_num
                elif 'produ' in s and 'manual' not in s and 'hora' not in s and 'produto' not in s and 'producao' not in col_map:
                    col_map['producao'] = col_num
                elif 'embalagem' in s and 'embalagem' not in col_map:
                    col_map['embalagem'] = col_num
                elif 'sleeve' in s and 'sleeve' not in col_map:
                    col_map['sleeve'] = col_num
                elif 'celofane' in s and 'celofane' not in col_map:
                    col_map['celofane'] = col_num
                elif 'turno' in s and 'turno' not in col_map:
                    col_map['turno'] = col_num
                elif 'colab' in s and 'total_colab' not in col_map:
                    col_map['total_colab'] = col_num
                elif 'ausente' in s and 'ausentes' not in col_map:
                    col_map['ausentes'] = col_num
            return r, col_map
    return None, {}


def main():
    print("=" * 70)
    print("INICIANDO MIGRAÇÃO HISTÓRICA DO EXCEL PARA O BANCO DE DADOS SQLITE")
    print("=" * 70)
    print(f"Diretório dos arquivos Excel: {EXCEL_FOLDER}")
    print(f"Banco de dados SQLite       : {DASHBOARD_DB}")

    # Ensure tables exist
    init_db()
    conn = get_db()
    cur = conn.cursor()

    inserted = {
        'producao_diaria': 0,
        'producao_manual': 0,
        'producao_manual_totais': 0,
        'observacoes_diarias': 0
    }

    # Tracking sets to avoid cross-file duplicates
    seen_envase = set()   # (ds, produto, prod_int)
    seen_emb = set()      # (ds, prod_int)
    seen_sleeve = set()   # (ds, prod_int)
    seen_cel = set()      # (ds, prod_int)
    seen_manual = set()   # (ds, produto, lote, fin, env)
    seen_obs = set()      # (ds, setor, obs_text)

    # Pre-populate seen sets from existing DB records
    existing_obs = cur.execute("SELECT data, setor, observacao FROM observacoes_diarias").fetchall()
    for r in existing_obs:
        seen_obs.add((r['data'], r['setor'], r['observacao']))

    existing_man = cur.execute("SELECT data, produto, coalesce(lote, ''), finalizados, envasados FROM producao_manual").fetchall()
    for r in existing_man:
        seen_manual.add((r[0], r[1], r[2], r[3], r[4]))

    # Locate source files: Produtividade_Geral.xlsx FIRST, then all daily files
    geral_file = os.path.join(EXCEL_FOLDER, 'Produtividade_Geral.xlsx')
    other_files = sorted(
        [f for f in glob.glob(os.path.join(EXCEL_FOLDER, 'Produtividade*.xlsx'))
         if not os.path.basename(f).startswith('~$')
         and 'BACKUP' not in f
         and os.path.abspath(f) != os.path.abspath(geral_file)],
        key=os.path.getmtime, reverse=True
    )

    file_list = []
    if os.path.exists(geral_file):
        file_list.append(geral_file)
    file_list.extend(other_files)

    print(f"\nArquivos encontrados para migração: {len(file_list)}")
    for i, fp in enumerate(file_list, 1):
        print(f"  {i:2d}. {os.path.basename(fp)}")

    # In-memory accumulators
    sector_prod = defaultdict(lambda: {'producao': 0, 'embalagem': 0, 'sleeve': 0, 'celofane': 0})
    manual_daily_totals = defaultdict(lambda: {'env': 0, 'fin': 0})

    # Process all files
    for fp in file_list:
        fname = os.path.basename(fp)
        try:
            wb = openpyxl.load_workbook(fp, data_only=True, read_only=True)
        except Exception as e:
            print(f"\n  [ERRO] Falha ao abrir {fname}: {e}")
            continue

        print(f"\n-> Processando: {fname}...")

        # 1) CONSOLIDADO TAB
        cons_name = None
        for s in wb.sheetnames:
            if 'consolidado' in s.lower() and 'antigo' not in s.lower():
                cons_name = s
                break

        if cons_name:
            ws = wb[cons_name]
            h_row, col_map = find_consolidado_cols(ws)
            if h_row and 'date' in col_map:
                d_col = col_map['date']
                p_col = col_map.get('producao')
                e_col = col_map.get('embalagem')
                s_col = col_map.get('sleeve')
                c_col = col_map.get('celofane')
                tc_col = col_map.get('total_colab')
                au_col = col_map.get('ausentes')

                cons_rows_count = 0
                for row in ws.iter_rows(min_row=h_row + 1, values_only=True):
                    if len(row) < d_col:
                        continue
                    dt_val = row[d_col - 1]
                    ds = parse_date(dt_val)
                    if not ds:
                        continue

                    prod = parse_int(row[p_col - 1]) if p_col and len(row) >= p_col else 0
                    emb = parse_int(row[e_col - 1]) if e_col and len(row) >= e_col else 0
                    slv = parse_int(row[s_col - 1]) if s_col and len(row) >= s_col else 0
                    cel = parse_int(row[c_col - 1]) if c_col and len(row) >= c_col else 0

                    tc = parse_int(row[tc_col - 1]) if tc_col and len(row) >= tc_col else None
                    au = parse_int(row[au_col - 1]) if au_col and len(row) >= au_col else None
                    if tc is not None and tc <= 0:
                        tc = None
                        au = None
                    elif tc is not None and au is None:
                        au = 0

                    if prod == 0 and emb == 0 and slv == 0 and cel == 0 and tc is None:
                        continue

                    turno = get_turno(ds)

                    cur.execute('''
                        INSERT OR IGNORE INTO producao_diaria 
                        (data, producao, embalagem, sleeve, celofane, turno, total_colab, ausentes)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                    ''', (ds, prod, emb, slv, cel, turno, tc, au))
                    if cur.rowcount > 0:
                        inserted['producao_diaria'] += 1
                        cons_rows_count += 1

                if cons_rows_count > 0:
                    print(f"   [Consolidado] +{cons_rows_count} registros diários inseridos")

        # 2) ENVASE TAB
        sname = find_sheet_ci(wb, 'Envase')
        if sname:
            ws = wb[sname]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or len(row) < 3:
                    continue
                dt_val = None
                prod_name = None
                p_val = None
                obs_val = None
                prob_val = None

                # Layout 1: [Data, Produto, Producao, OBS, Prob. Maquinario] (Geral)
                if parse_date(row[0]):
                    dt_val = row[0]
                    prod_name = row[1] if len(row) > 1 else ''
                    p_val = row[2] if len(row) > 2 else 0
                    obs_val = row[3] if len(row) > 3 else None
                    prob_val = row[4] if len(row) > 4 else None
                # Layout 2: [DIA, DATA, PRODUTO, Producao, OBS, Prob. Maquinario] (Daily files)
                elif len(row) > 1 and parse_date(row[1]):
                    dt_val = row[1]
                    prod_name = row[2] if len(row) > 2 else ''
                    p_val = row[3] if len(row) > 3 else (row[2] if len(row) > 2 else 0)
                    obs_val = row[4] if len(row) > 4 else None
                    prob_val = row[5] if len(row) > 5 else None

                ds = parse_date(dt_val)
                if not ds:
                    continue

                prod = parse_int(p_val)
                prod_str = str(prod_name or '').strip()
                k = (ds, prod_str, prod)
                if k not in seen_envase:
                    seen_envase.add(k)
                    if prod > 0:
                        sector_prod[ds]['producao'] += prod

                # Observations
                obs1 = clean_obs(obs_val)
                obs2 = clean_obs(prob_val)
                combined = []
                if obs1:
                    combined.append(obs1)
                if obs2 and obs2 != obs1:
                    combined.append(obs2)
                if combined:
                    full_obs = ' | '.join(combined)
                    obs_k = (ds, 'Envase', full_obs)
                    if obs_k not in seen_obs:
                        seen_obs.add(obs_k)
                        cur.execute('''
                            INSERT OR IGNORE INTO observacoes_diarias (data, setor, observacao)
                            VALUES (?, ?, ?)
                        ''', (ds, 'Envase', full_obs))
                        if cur.rowcount > 0:
                            inserted['observacoes_diarias'] += 1

        # 3) EMBALAGEM TAB
        sname = find_sheet_ci(wb, 'Embalagem')
        if sname:
            ws = wb[sname]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or len(row) < 2:
                    continue
                dt_val = None
                p_val = None
                obs_val = None

                if parse_date(row[0]):
                    dt_val = row[0]
                    p_val = row[1] if len(row) > 1 else 0
                    obs_val = row[2] if len(row) > 2 else None
                elif len(row) > 1 and parse_date(row[1]):
                    dt_val = row[1]
                    p_val = row[2] if len(row) > 2 else 0
                    obs_val = row[3] if len(row) > 3 else None

                ds = parse_date(dt_val)
                if not ds:
                    continue

                prod = parse_int(p_val)
                k = (ds, prod)
                if k not in seen_emb:
                    seen_emb.add(k)
                    if prod > 0:
                        sector_prod[ds]['embalagem'] += prod

                obs = clean_obs(obs_val)
                if obs:
                    obs_k = (ds, 'Embalagem', obs)
                    if obs_k not in seen_obs:
                        seen_obs.add(obs_k)
                        cur.execute('''
                            INSERT OR IGNORE INTO observacoes_diarias (data, setor, observacao)
                            VALUES (?, ?, ?)
                        ''', (ds, 'Embalagem', obs))
                        if cur.rowcount > 0:
                            inserted['observacoes_diarias'] += 1

        # 4) SLEEVE TAB
        sname = find_sheet_ci(wb, 'Sleeve')
        if sname:
            ws = wb[sname]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or len(row) < 2:
                    continue
                dt_val = None
                p_val = None
                obs_val = None

                if parse_date(row[0]):
                    dt_val = row[0]
                    p_val = row[1] if len(row) > 1 else 0
                    obs_val = row[2] if len(row) > 2 else None
                elif len(row) > 1 and parse_date(row[1]):
                    dt_val = row[1]
                    p_val = row[2] if len(row) > 2 else 0
                    obs_val = row[3] if len(row) > 3 else None

                ds = parse_date(dt_val)
                if not ds:
                    continue

                prod = parse_int(p_val)
                k = (ds, prod)
                if k not in seen_sleeve:
                    seen_sleeve.add(k)
                    if prod > 0:
                        sector_prod[ds]['sleeve'] += prod

                obs = clean_obs(obs_val)
                if obs:
                    obs_k = (ds, 'Sleeve', obs)
                    if obs_k not in seen_obs:
                        seen_obs.add(obs_k)
                        cur.execute('''
                            INSERT OR IGNORE INTO observacoes_diarias (data, setor, observacao)
                            VALUES (?, ?, ?)
                        ''', (ds, 'Sleeve', obs))
                        if cur.rowcount > 0:
                            inserted['observacoes_diarias'] += 1

        # 5) CELOFANE TAB
        sname = find_sheet_ci(wb, 'Celofane')
        if sname:
            ws = wb[sname]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or len(row) < 2:
                    continue
                dt_val = None
                p_val = None
                obs_val = None

                if parse_date(row[0]):
                    dt_val = row[0]
                    p_val = row[1] if len(row) > 1 else 0
                    obs_val = row[2] if len(row) > 2 else None
                elif len(row) > 1 and parse_date(row[1]):
                    dt_val = row[1]
                    p_val = row[2] if len(row) > 2 else 0
                    obs_val = row[3] if len(row) > 3 else None

                ds = parse_date(dt_val)
                if not ds:
                    continue

                prod = parse_int(p_val)
                k = (ds, prod)
                if k not in seen_cel:
                    seen_cel.add(k)
                    if prod > 0:
                        sector_prod[ds]['celofane'] += prod

                obs = clean_obs(obs_val)
                if obs:
                    obs_k = (ds, 'Celofane', obs)
                    if obs_k not in seen_obs:
                        seen_obs.add(obs_k)
                        cur.execute('''
                            INSERT OR IGNORE INTO observacoes_diarias (data, setor, observacao)
                            VALUES (?, ?, ?)
                        ''', (ds, 'Celofane', obs))
                        if cur.rowcount > 0:
                            inserted['observacoes_diarias'] += 1

        # 6) ENVASE MANUAL TAB
        sname = find_sheet_ci(wb, 'Envase Manual')
        if sname:
            ws = wb[sname]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or len(row) < 4:
                    continue
                ds = parse_date(row[0])
                if not ds:
                    continue
                produto = str(row[1] or '').strip()
                lote = str(row[2] or '').strip() if len(row) > 2 and row[2] else None
                fin = parse_int(row[3]) if len(row) > 3 else 0
                env = parse_int(row[4]) if len(row) > 4 else 0
                ret = parse_int(row[5]) if len(row) > 5 else 0
                obs = clean_obs(row[6]) if len(row) > 6 else None

                if not produto or (fin == 0 and env == 0):
                    continue

                mk = (ds, produto, lote or '', fin, env)
                if mk not in seen_manual:
                    seen_manual.add(mk)
                    cur.execute('''
                        INSERT INTO producao_manual (data, produto, lote, envasados, finalizados, retrabalhadas)
                        VALUES (?, ?, ?, ?, ?, ?)
                    ''', (ds, produto, lote, env, fin, ret))
                    if cur.rowcount > 0:
                        inserted['producao_manual'] += 1

                manual_daily_totals[ds]['env'] += env
                manual_daily_totals[ds]['fin'] += fin

                if obs:
                    obs_k = (ds, 'Envase Manual', obs)
                    if obs_k not in seen_obs:
                        seen_obs.add(obs_k)
                        cur.execute('''
                            INSERT OR IGNORE INTO observacoes_diarias (data, setor, observacao)
                            VALUES (?, ?, ?)
                        ''', (ds, 'Envase Manual', obs))
                        if cur.rowcount > 0:
                            inserted['observacoes_diarias'] += 1

        wb.close()

    # Insert remaining historical production dates from aggregated sectors (2024, 2025, etc.)
    print("\nInserindo datas de produção histórica agregadas dos setores...")
    for ds, counts in sorted(sector_prod.items()):
        p = counts['producao']
        e = counts['embalagem']
        s = counts['sleeve']
        c = counts['celofane']
        if p == 0 and e == 0 and s == 0 and c == 0:
            continue
        turno = get_turno(ds)
        cur.execute('''
            INSERT OR IGNORE INTO producao_diaria 
            (data, producao, embalagem, sleeve, celofane, turno, total_colab, ausentes)
            VALUES (?, ?, ?, ?, ?, ?, NULL, NULL)
        ''', (ds, p, e, s, c, turno))
        if cur.rowcount > 0:
            inserted['producao_diaria'] += 1

    # Insert daily totals for manual packaging
    print("Inserindo totais diários de envase manual...")
    for ds, totals in sorted(manual_daily_totals.items()):
        cur.execute('''
            INSERT OR IGNORE INTO producao_manual_totais (data, total_envasados, total_finalizados)
            VALUES (?, ?, ?)
        ''', (ds, totals['env'], totals['fin']))
        if cur.rowcount > 0:
            inserted['producao_manual_totais'] += 1

    # Commit all changes
    conn.commit()

    # Final summary report
    print("\n" + "=" * 70)
    print("MIGRAÇÃO CONCLUÍDA COM SUCESSO!")
    print("=" * 70)
    print("Registros recém-importados nesta execução:")
    for tbl, count in inserted.items():
        print(f"  + {tbl:24}: {count} novos registros")

    print("\nTotal acumulado de registros no banco de dados SQLite:")
    for tbl in [
        'producao_diaria',
        'producao_manual',
        'producao_manual_totais',
        'observacoes_diarias',
        'metas_periodo',
        'metas_diarias',
        'absenteismo',
        'configuracoes'
    ]:
        try:
            total = cur.execute(f"SELECT count(*) FROM {tbl}").fetchone()[0]
            print(f"  - {tbl:24}: {total} registros")
        except Exception:
            pass

    min_d, max_d = cur.execute("SELECT min(data), max(data) FROM producao_diaria").fetchone()
    print(f"\nPeríodo histórico coberto em producao_diaria: {min_d} a {max_d}")
    print("=" * 70)

    conn.close()


if __name__ == '__main__':
    main()
