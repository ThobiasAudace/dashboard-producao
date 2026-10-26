"""
Manual Line data source.

Reads production data from the dashboard's own database (producao_manual tables).
Includes a WhatsApp message parser to extract production data automatically.

Future: can also read from sistema_producao/producao.db as a secondary source.
"""

import os
import re
from datetime import datetime, timedelta, date

from database import get_db
from config import META_LINHA_MANUAL


DIAS_SEMANA = {0: 'SEG', 1: 'TER', 2: 'QUA', 3: 'QUI', 4: 'SEX', 5: 'SÁB', 6: 'DOM'}


# ========== WhatsApp Message Parser ==========

def parse_whatsapp_message(texto):
    """Parse a standardized WhatsApp message for manual line production.

    Handles multiple formats:
        Envase Manual-                          (simple header)
        Envase manual dia 01/09/2026:           (header with date)

        My soul breeze L.1201/26:               (product with lot)
        Retrabalho shower gel:                  (product without lot)

        Envasados: 1.576 unidades
        Finalizados: 2.466 unidades
        Total retrabalhadas: 560 unidades

        Total envasados: 3.093 unidades         (full form)
        Total En: 3.171 unidades                (abbreviated)
        Total Fi: 6.105 unidades                (abbreviated)
        Total geral finalizados 30 ml: 3.923    (total geral)

    Returns:
        dict with keys: produtos (list), total_envasados, total_finalizados,
                        total_retrabalhadas, data_extraida
    """
    produtos = []
    total_envasados = 0
    total_finalizados = 0
    total_retrabalhadas = 0
    data_extraida = None

    # Normalize text
    texto = texto.strip()
    lines = [l.strip() for l in texto.split('\n') if l.strip()]

    current_produto = None
    current_lote = None
    current_env = 0
    current_fin = 0
    current_retrab = 0
    observacao = None

    def _save_current():
        """Save the current product to the list."""
        if current_produto:
            produtos.append({
                'produto': current_produto,
                'lote': current_lote or '',
                'envasados': current_env,
                'finalizados': current_fin,
                'retrabalhadas': current_retrab,
            })

    for line in lines:
        # Skip/parse header line (may contain date)
        if re.match(r'envase\s+manual', line, re.IGNORECASE):
            date_match = re.search(r'(\d{1,2})[/\-](\d{1,2})[/\-](\d{4})', line)
            if date_match:
                try:
                    data_extraida = f"{date_match.group(3)}-{int(date_match.group(2)):02d}-{int(date_match.group(1)):02d}"
                except (ValueError, TypeError):
                    pass
            continue

        # ---- TOTAL LINES (check before product lines) ----

        # Total retrabalhadas / Total retrab
        match_total_retrab = re.match(r'total\s+retrab\w*\s*:\s*([\d.]+)', line, re.IGNORECASE)
        if match_total_retrab:
            val = _parse_number(match_total_retrab.group(1))
            # Could be under a product or a standalone total
            if current_produto and current_retrab == 0:
                current_retrab = val
            else:
                total_retrabalhadas += val
            continue

        # Total geral finalizados / Total geral fin
        match_total_geral_fin = re.match(r'total\s+geral\s+fi\w*.*?:\s*([\d.]+)', line, re.IGNORECASE)
        if match_total_geral_fin:
            total_finalizados = _parse_number(match_total_geral_fin.group(1))
            continue

        # Total geral envasados / Total geral en
        match_total_geral_env = re.match(r'total\s+geral\s+en\w*.*?:\s*([\d.]+)', line, re.IGNORECASE)
        if match_total_geral_env:
            total_envasados = _parse_number(match_total_geral_env.group(1))
            continue

        # Total envasados / Total En / Total Env
        match_total_env = re.match(r'total\s+en\w*\s*:\s*([\d.]+)', line, re.IGNORECASE)
        if match_total_env:
            total_envasados = _parse_number(match_total_env.group(1))
            continue

        # Total finalizados / Total Fi / Total Fin
        match_total_fin = re.match(r'total\s+fi\w*\s*:\s*([\d.]+)', line, re.IGNORECASE)
        if match_total_fin:
            total_finalizados = _parse_number(match_total_fin.group(1))
            continue

        # ---- OBSERVATION LINES ----
        match_obs = re.match(r'^obs(?:erva[çc][aã]o)?s?\s*:\s*(.+)', line, re.IGNORECASE)
        if match_obs:
            observacao = match_obs.group(1).strip()
            continue

        # ---- PRODUCT LINES ----

        # Product with lot number (L.XXXX/XX)
        match_produto_lote = re.match(r'(.+?)\s+(L\.?\s*\d+/\d+)\s*:?\s*$', line, re.IGNORECASE)
        if match_produto_lote:
            _save_current()
            current_produto = match_produto_lote.group(1).strip()
            current_lote = match_produto_lote.group(2).strip()
            current_env = 0
            current_fin = 0
            current_retrab = 0
            continue

        # Product without lot (e.g. "Retrabalho shower gel:")
        # Must end with colon, and NOT match value patterns
        match_produto_sem_lote = re.match(r'^([A-Za-zÀ-ú][\w\sÀ-ú]+?)\s*:\s*$', line)
        if match_produto_sem_lote:
            nome = match_produto_sem_lote.group(1).strip()
            # Skip if it looks like a value line
            if not re.match(r'(envasados?|finalizados?|total)', nome, re.IGNORECASE):
                _save_current()
                current_produto = nome
                current_lote = None
                current_env = 0
                current_fin = 0
                current_retrab = 0
                continue

        # ---- VALUE LINES ----

        # Retrabalhadas
        match_retrab = re.match(r'retrabalhad\w*\s*:\s*([\d.]+)', line, re.IGNORECASE)
        if match_retrab:
            current_retrab = _parse_number(match_retrab.group(1))
            continue

        # Envasados
        match_env = re.match(r'envasados?\s*:\s*([\d.]+)', line, re.IGNORECASE)
        if match_env:
            current_env = _parse_number(match_env.group(1))
            continue

        # Finalizados
        match_fin = re.match(r'finalizados?\s*:\s*([\d.]+)', line, re.IGNORECASE)
        if match_fin:
            current_fin = _parse_number(match_fin.group(1))
            continue

        # Bare number (e.g. "178 unidades" without a field prefix)
        # Assign based on context: retrabalho -> retrabalhadas, else -> envasados
        match_bare = re.match(r'^([\d.]+)\s*(?:unidades?|un)?\s*$', line, re.IGNORECASE)
        if match_bare and current_produto:
            val = _parse_number(match_bare.group(1))
            if current_produto and re.search(r'retrabalho', current_produto, re.IGNORECASE):
                current_retrab = val
            else:
                current_env = val
            continue

    # Save last product
    _save_current()

    # Calculate totals from products if not in message
    if total_envasados == 0 and produtos:
        total_envasados = sum(p['envasados'] for p in produtos)
    if total_finalizados == 0 and produtos:
        total_finalizados = sum(p['finalizados'] for p in produtos)
    if total_retrabalhadas == 0 and produtos:
        total_retrabalhadas = sum(p.get('retrabalhadas', 0) for p in produtos)

    return {
        'produtos': produtos,
        'total_envasados': total_envasados,
        'total_finalizados': total_finalizados,
        'total_retrabalhadas': total_retrabalhadas,
        'data_extraida': data_extraida,
        'observacao': observacao,
    }


def _parse_number(text):
    """Parse Brazilian number format: 1.576 -> 1576, 3.1058 -> 31058"""
    return int(text.replace('.', '').replace(',', '').strip())


# ========== Database Operations ==========

def salvar_producao_manual(data, produtos, total_envasados, total_finalizados, observacao='', total_retrabalhadas=0):
    """Save manual line production data from parsed WhatsApp message."""
    conn = get_db()
    cursor = conn.cursor()

    # Add retrabalhadas column if not exists (migration)
    try:
        cursor.execute("ALTER TABLE producao_manual ADD COLUMN retrabalhadas INTEGER DEFAULT 0")
    except Exception:
        pass  # Column already exists

    # Clear existing records for this date
    cursor.execute("DELETE FROM producao_manual WHERE data = ?", (data,))

    # Insert product details
    for p in produtos:
        cursor.execute("""
            INSERT INTO producao_manual (data, produto, lote, envasados, finalizados, retrabalhadas)
            VALUES (?, ?, ?, ?, ?, ?)
        """, (data, p['produto'], p.get('lote', ''), p['envasados'], p['finalizados'], p.get('retrabalhadas', 0)))

    # Insert/update daily totals
    cursor.execute("""
        INSERT OR REPLACE INTO producao_manual_totais
            (data, total_envasados, total_finalizados, observacao)
        VALUES (?, ?, ?, ?)
    """, (data, total_envasados, total_finalizados, observacao))

    conn.commit()
    conn.close()

    # Sync to Excel spreadsheet
    _sync_to_excel(data, produtos, total_retrabalhadas)


def _sync_to_excel(data, produtos, total_retrabalhadas=0):
    """Sync manual line data to the 'Envase Manual' tab in the Excel file."""
    import openpyxl
    from config import EXCEL_FOLDER

    # Find the latest Produtividade file
    import glob
    pattern = os.path.join(EXCEL_FOLDER, 'Produtividade*.xlsx')
    files = [f for f in glob.glob(pattern) if not os.path.basename(f).startswith('~$')]
    if not files:
        print("[manual_line] No Excel file found for sync")
        return
    files.sort(key=os.path.getmtime, reverse=True)

    for file_path in files:
        try:
            wb = openpyxl.load_workbook(file_path)
            if 'Envase Manual' not in wb.sheetnames:
                print(f"[manual_line] Tab 'Envase Manual' not found in {os.path.basename(file_path)}")
                wb.close()
                continue

            ws = wb['Envase Manual']

            # Remove existing rows for this date
            rows_to_delete = []
            for row in range(2, ws.max_row + 1):
                cell_val = ws.cell(row, 1).value
                if cell_val is None:
                    continue
                cell_str = cell_val
                if hasattr(cell_val, 'strftime'):
                    cell_str = cell_val.strftime('%Y-%m-%d')
                if str(cell_str) == data:
                    rows_to_delete.append(row)

            for row in sorted(rows_to_delete, reverse=True):
                ws.delete_rows(row)

            # Find the right position to insert (keep sorted by date)
            insert_row = ws.max_row + 1
            for row in range(2, ws.max_row + 1):
                cell_val = ws.cell(row, 1).value
                if cell_val is None:
                    insert_row = row
                    break
                cell_str = cell_val
                if hasattr(cell_val, 'strftime'):
                    cell_str = cell_val.strftime('%Y-%m-%d')
                if str(cell_str) > data:
                    insert_row = row
                    break

            # Insert new rows for each product
            from datetime import datetime as dt
            data_date = dt.strptime(data, '%Y-%m-%d').date()

            for i, p in enumerate(produtos):
                ws.insert_rows(insert_row + i)
                ws.cell(insert_row + i, 1, data_date).number_format = 'DD/MM/YYYY'
                ws.cell(insert_row + i, 2, p['produto'])
                ws.cell(insert_row + i, 3, p.get('lote', ''))
                ws.cell(insert_row + i, 4, p.get('finalizados', 0))
                ws.cell(insert_row + i, 5, p.get('envasados', 0))
                ws.cell(insert_row + i, 6, p.get('retrabalhadas', 0))
                ws.cell(insert_row + i, 7, '')  # OBS — user fills manually

            wb.save(file_path)
            wb.close()
            print(f"[manual_line] Synced {len(produtos)} products to Excel for {data}")
            return

        except PermissionError:
            print(f"[manual_line] Excel locked: {os.path.basename(file_path)}, skip sync")
            continue
        except Exception as e:
            print(f"[manual_line] Excel sync error: {e}")
            continue


def get_produtos_dia(data):
    """Get product-level detail for a specific date."""
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute(
        "SELECT * FROM producao_manual WHERE data = ? ORDER BY id",
        (data,)
    )
    rows = cursor.fetchall()
    conn.close()
    return [dict(r) for r in rows]


def deletar_producao_dia(data):
    """Delete all manual line records for a date."""
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM producao_manual WHERE data = ?", (data,))
    cursor.execute("DELETE FROM producao_manual_totais WHERE data = ?", (data,))
    conn.commit()
    conn.close()


# ========== Dashboard Data Functions ==========

def get_dados_dia(data=None):
    """Get manual line production data for a specific date."""
    if data is None:
        data = date.today().strftime('%Y-%m-%d')
    elif isinstance(data, (date, datetime)):
        data = data.strftime('%Y-%m-%d')

    conn = get_db()
    cursor = conn.cursor()

    # Read from local dashboard database
    cursor.execute(
        "SELECT * FROM producao_manual_totais WHERE data = ?",
        (data,)
    )
    row = cursor.fetchone()

    if row:
        envasados = row['total_envasados'] or 0
        finalizados = row['total_finalizados'] or 0
    else:
        envasados = 0
        finalizados = 0

    conn.close()

    from config import get_meta_periodo
    d_meta = get_meta_periodo(int(data[:4]), int(data[5:7]), data_str=data)
    meta = d_meta['manual']
    eficiencia = (finalizados / meta * 100) if meta > 0 else 0

    return {
        'data': data,
        'envasados': envasados,
        'finalizados': finalizados,
        'total': envasados + finalizados,
        'meta': meta,
        'eficiencia': round(eficiencia, 1),
    }


def get_dados_semana(data_ref=None, dias=7):
    """Get manual line data for the last N days."""
    if data_ref is None:
        data_ref = date.today()
    elif isinstance(data_ref, str):
        data_ref = datetime.strptime(data_ref, '%Y-%m-%d').date()

    start = data_ref - timedelta(days=dias - 1)
    result = []

    for i in range(dias):
        day = start + timedelta(days=i)
        day_str = day.strftime('%Y-%m-%d')
        day_data = get_dados_dia(day_str)
        day_data['dia_semana'] = DIAS_SEMANA.get(day.weekday(), '')
        result.append(day_data)

    return result


def get_dados_mes(ano=None, mes=None):
    """Get manual line data for an entire month."""
    if ano is None or mes is None:
        hoje = date.today()
        ano = hoje.year
        mes = hoje.month

    conn = get_db()
    cursor = conn.cursor()
    data_inicio = f"{ano}-{mes:02d}-01"

    if mes == 12:
        data_fim = f"{ano + 1}-01-01"
    else:
        data_fim = f"{ano}-{mes + 1:02d}-01"

    cursor.execute("""
        SELECT data, total_envasados, total_finalizados
        FROM producao_manual_totais
        WHERE data >= ? AND data < ?
        ORDER BY data
    """, (data_inicio, data_fim))
    rows = cursor.fetchall()
    conn.close()

    from config import get_meta_periodo
    result = []
    for row in rows:
        env = row['total_envasados'] or 0
        fin = row['total_finalizados'] or 0
        d_str = row['data']
        d_meta = get_meta_periodo(int(d_str[:4]), int(d_str[5:7]), data_str=d_str)
        meta = d_meta['manual']
        eficiencia = (fin / meta * 100) if meta > 0 else 0

        result.append({
            'data': d_str,
            'envasados': env,
            'finalizados': fin,
            'total': env + fin,
            'meta': meta,
            'eficiencia': round(eficiencia, 1),
        })

    return result


def get_resumo_mes(ano=None, mes=None):
    """Get monthly summary for manual line."""
    dados = get_dados_mes(ano, mes)
    if not dados:
        return {
            'total_envasados': 0, 'total_finalizados': 0,
            'total_meta': 0, 'eficiencia_media': 0, 'dias_trabalhados': 0
        }

    total_env = sum(d['envasados'] for d in dados)
    total_fin = sum(d['finalizados'] for d in dados)
    dias = len([d for d in dados if d['envasados'] > 0 or d['finalizados'] > 0])
    total_meta = sum(d['meta'] for d in dados if d['envasados'] > 0 or d['finalizados'] > 0)
    eficiencia = (total_fin / total_meta * 100) if total_meta > 0 else 0

    return {
        'total_envasados': total_env,
        'total_finalizados': total_fin,
        'total_meta': total_meta,
        'eficiencia_media': round(eficiencia, 1),
        'dias_trabalhados': dias,
    }
